"""Read-only verification of key manuscript numbers against new computations."""
from pathlib import Path
import json,hashlib,sys
import numpy as np,pandas as pd
from docx import Document
ROOT=Path(__file__).resolve().parents[2]
rd=lambda p:pd.read_csv(ROOT/p,sep='\t',dtype={'sample_id':str})
main=Document(ROOT/'reference/01_论文正文.docx')
supp=Document(ROOT/'reference/04_补充材料.docx')
checks=[]
def ck(label,actual,expected,source,atol=0,rtol=0):
    if isinstance(expected,str):ok=str(actual)==expected
    else:ok=bool(np.allclose(actual,expected,atol=atol,rtol=rtol,equal_nan=True))
    checks.append(dict(claim=label,actual=actual,manuscript=expected,passed=ok,source=source))
def rounded(values,places=3):return [float(round(float(x),places)) for x in values]
def sig(values,precision=3):return [float(f'{float(x):.{precision}g}') for x in values]
def upper(label,value,bound,source):
    checks.append(dict(claim=label,actual=float(value),manuscript='<= '+str(bound),passed=bool(value<=bound),source=source))
prep='runs/preparation/project/'
raw=rd(prep+'03_分析与结果/agp_models/raw_covariates.tsv')
base=raw[raw.total_reads_qc.ge(5000)].copy()
ck('描述样本2760',len(base),2760,'fresh raw_covariates with total_reads_qc>=5000')
table={r.cells[0].text:r.cells[1].text for r in main.tables[0].rows[1:]}
def summarize(x):
    q=x.quantile([.5,.25,.75]);return f'{q.iloc[0]:.1f}（{q.iloc[1]:.1f}–{q.iloc[2]:.1f}）'
for label,x in [('年龄（岁）',base.age),('BMI（kg/m²）',base.bmi),('总测序读数',base.total_reads_qc),('纤维密度（g/MJ）',base.fiber_density/4.184)]:
    ck('表1 '+label,summarize(x),table[label],'fresh source-derived raw_covariates; fiber g/1000 kcal divided by 4.184')
cat=[('女性',base.sex.eq('female')),('男性',base.sex.eq('male')),('其他或未说明性别',base.sex.eq('other_unknown')),('美国',base.country.eq('US')),('英国',base.country.eq('UK')),('其他地区',base.country.eq('Other')),('地区未提供',base.country.eq('Unknown')),('过去1年抗生素使用',base.antibiotic.eq(1)),('炎症性肠病',base.ibd.eq(1)),('肠易激综合征',base.ibs.eq(1))]
for label,mask in cat:
    n=int(mask.sum());ck('表1 '+label,f'{n}（{n/len(base)*100:.1f}%）',table[label],'fresh source-derived raw_covariates')
ck('表1脚注BMI缺失',int(base.bmi.isna().sum()),35,'fresh raw_covariates')
ck('表1其余连续变量无缺失',int(base[['age','total_reads_qc','fiber_density']].isna().sum().sum()),0,'fresh raw_covariates')
Y=rd(prep+'03_分析与结果/association_models/association_target_counts.tsv').set_index('sample_id')
allY=rd(prep+'05_输入数据/派生输入/lactobacillales_genus_counts_all_samples.tsv').set_index('sample_id')
ck('主模型样本',len(Y),2748,'fresh target counts')
ck('主模型六属检出人数',Y.ge(1).sum().tolist(),[502,469,191,158,147,90],'fresh target counts')
mainY=allY.loc[base.sample_id,Y.columns]
ck('描述样本乳酪杆菌/乳杆菌/片球菌检出人数',mainY[['Lacticaseibacillus','Lactobacillus','Pediococcus']].ge(1).sum().tolist(),[504,469,90],'fresh all-sample counts, 2760 base sample')
ck('描述样本检出率范围',rounded([mainY.ge(1).mean().min()*100,mainY.ge(1).mean().max()*100],2),[3.26,18.26],'fresh all-sample counts')
om=rd('runs/core/refitted_primary_pc_omnibus.tsv').iloc[0]
ck('六属整体chi2',round(float(om.chi2),2),174.09,'new core omnibus fit')
ck('六属整体df',int(om.df),50,'new core omnibus fit')
ck('六属整体P',sig([om.p]),[1.25e-15],'new core omnibus fit')
c1=rd('runs/core/refitted_936_associations.tsv');qcol=next(c for c in c1 if c.startswith('q'))
ck('936项检验数',len(c1),936,'new 936 regression fits')
ck('936项校正显著数及背景属数',[int(c1[qcol].lt(.05).sum()),int(c1.loc[c1[qcol].lt(.05),'feature'].nunique())],[120,73],'new 936 regression fits')
c4=rd('runs/core/表S22c_全部15组主成分向量比较.tsv')
ck('15项比较BH及Holm显著数',[int(c4.q_BH_15.lt(.05).sum()),int(c4.p_Holm_15.lt(.05).sum())],[12,9],'new core pairwise PC fits')
row=c4.iloc[0]
ck('最强目标配对',row.taxon_A+'/'+row.taxon_B,'Lacticaseibacillus/Lactobacillus','new core pairwise PC fits')
ck('最强配对chi2',round(row.chi2,2),88.14,'new core pairwise PC fits')
ck('最强配对q',sig([row.q_BH_15]),[1.87e-13],'new core pairwise PC fits')
c5=rd('runs/core/表S22e_全部156背景属整体差异.tsv');qc=next(c for c in c5 if c.startswith('q'))
c6=rd('runs/core/表S22f_全部2340项背景属系数比较.tsv')
ck('背景整体及2340项显著数',[int(c5[qc].lt(.05).sum()),int(c6.q_BH_2340.lt(.05).sum())],[35,71],'new core background comparisons')
BG=['Anaerococcus','Finegoldia','Peptoniphilus_A']
effects=rd('runs/core/表S24i_三背景两目标属效应量.tsv')
for target,expected in [('Lactobacillus',[1.482,1.519,1.429]),('Lacticaseibacillus',[.955,1.003,.938])]:
    t=effects[(effects.scenario=='original_reads1')&(effects.taxon==target)]
    ck('基础OR '+target,rounded(t.OR),expected,'new C13')
top=c6.nsmallest(3,'q_BH_2340')
ck('最小三项q',sig(top.q_BH_2340),[7.58e-8,7.58e-8,1.16e-7],'new C06')
p1=rd('runs/deepening/results/P01_广泛饮食调整全部关联.tsv');p2=rd('runs/deepening/results/P02_广泛饮食调整全部差异.tsv')
for target,expected in [('Lactobacillus',[1.485,1.520,1.432]),('Lacticaseibacillus',[.945,.989,.934])]:
    t=p1[(p1.spec=='food10_nutrient10')&(p1.target==target)]
    ck('饮食调整OR '+target,rounded(t.OR),expected,'new P01')
t=p2[p2.spec=='food10_nutrient10']
for col,expected in [('ratio_OR',[.637,.651,.652]),('low',[.557,.573,.572]),('high',[.728,.739,.744])]:ck('饮食调整OR比 '+col,rounded(t[col]),expected,'new P02')
ps=json.loads((ROOT/'runs/deepening/diet_preprocessing/饮食预检验汇总.json').read_text(encoding='utf8'))
ck('食品营养维数',[ps['preprocessing'][v]['n_features'] for v in ['food_PC','nutrient_PC']],[177,184],'new preprocessing')
ck('食品营养10PC解释变异%',rounded([ps['preprocessing'][v]['explained_variance_cumulative']['10']*100 for v in ['food_PC','nutrient_PC']],2),[25.96,73.84],'new PCA')
ck('能量限制排除26人',ps['energy_restriction']['excluded'],26,'new energy filter')
z=rd(prep+'03_分析与结果/association_models/association_standardized_clr.tsv').set_index('sample_id');corr=z[['microbe::Peptoniphilaceae|'+g for g in BG]].corr(method='spearman').to_numpy();vals=corr[np.triu_indices(3,1)]
ck('三背景Spearman范围',rounded([vals.min(),vals.max()]),[.688,.713],'fresh CLR')
j3=rd('runs/deepening/results/J03_联合整体差异.tsv');t=j3[(j3.threshold==1)&(j3['mode']=='three_joint')].iloc[0]
ck('三背景联合chi2',round(t.chi2,2),49.67,'new J03');ck('三背景联合P',sig([t.p]),[9.40e-11],'new J03')
j1=rd('runs/deepening/results/J01_联合与菌科关联.tsv');t=j1[(j1.threshold==1)&(j1['mode']=='whole_family')&(j1.target=='Lactobacillus')].iloc[0]
ck('菌科OR及CI',rounded([t.OR,t.low,t.high]),[1.549,1.383,1.734],'new J01')
j2=rd('runs/deepening/results/J02_联合与菌科直接差异.tsv');t=j2[(j2.threshold==1)&(j2['mode']=='whole_family')].iloc[0]
ck('菌科OR比及CI',rounded([t.ratio_OR,t.low,t.high]),[.651,.560,.757],'new J02')
ck('联合单项无显著数',int(j2[(j2.threshold==1)&(j2['mode']=='three_joint')].q_BH.lt(.05).sum()),0,'new J02')
state=Y.Lacticaseibacillus.ge(1).astype(int)+2*Y.Lactobacillus.ge(1).astype(int)
ck('四种状态人数',state.value_counts().sort_index().tolist(),[1926,353,320,149],'fresh target counts')
s1=rd('runs/deepening/results/S01_四种检出状态直接比较.tsv')
for stateid,expected in [(2,[1.616,1.567,1.582]),(3,[1.377,1.317,1.394])]:
    t=s1[(s1.threshold==1)&s1.background.isin(BG)&(s1.state_A==stateid)&(s1.state_B==1)]
    ck('状态RRR '+str(stateid),rounded(t.RRR),expected,'new S01');ck('状态对比全显著 '+str(stateid),int(t.q_BH.lt(.05).sum()),3,'new S01')
s2=rd('runs/deepening/results/S02_四状态标准化概率.tsv');t=s2[(s2.threshold==1)&s2.background.isin(BG)&s2.state.isin([2,3])].pivot(index=['background','state'],columns='level',values='probability')
ck('乳杆菌单独/共同检出高背景概率更高',int((t.p90>t.p10).sum()),6,'new S02')
a1=rd('runs/deepening/results/A01_检出后相对丰度.tsv');t=a1[(a1.threshold==1)&(a1.target=='Lactobacillus')]
for col,expected in [('ratio_geomean',[1.427,1.451,1.406]),('low',[1.233,1.265,1.233]),('high',[1.652,1.664,1.603])]:ck('正丰度 '+col,rounded(t[col]),expected,'new A01')
upper('乳杆菌丰度q上限',t.q_BH_6.max(),3.81e-6,'new A01')
t=a1[(a1.threshold==1)&(a1.target=='Lacticaseibacillus')];ck('乳酪杆菌丰度比',rounded(t.ratio_geomean),[.925,.957,1.027],'new A01');ck('乳酪杆菌丰度显著数',int(t.q_BH_6.lt(.05).sum()),0,'new A01')
ck('≥3阈值丰度实际n',a1[a1.threshold==3].groupby('target').n.first().loc[['Lacticaseibacillus','Lactobacillus']].tolist(),[380,343],'new A01')
a2=rd('runs/deepening/sex_results/A02_最高读数样本剔除复核.tsv');t=a2[a2.target=='Lactobacillus'];ck('剔除高读数丰度比',rounded(t.ratio_geomean),[1.399,1.423,1.391],'new A02');ck('最高读数占比%',round(float(t.fraction_of_target_reads.iloc[0])*100,2),75.89,'new A02')
ck('高读数剔除后n',int(t.n.iloc[0]),468,'new A02');upper('高读数剔除后q上限',t.q_BH.max(),8.66e-6,'new A02')
h=rd('runs/deepening/sex_results/H02_属间差异及性别直接比较.tsv');t=h[(h.threshold==1)&(h.scheme=='primary')&(h.comparison=='sex_difference')]
ck('性别差异比',rounded(t.ratio_OR),[1.485,1.454,1.250],'new H02');ck('性别差异q',sig(t.q_BH_3),[.00625,.00625,.0964],'new H02')
v=rd('runs/deepening/results/V01_ASV候选完整列表.tsv');ck('三背景ASV数',v.groupby('genus').size().reindex(BG).tolist(),[3,1,3],'new V01');ck('乳杆菌两条ASV检出人数',v[v.genus=='Lactobacillus'].positive_people.tolist(),[261,144],'new V01')
l=rd('runs/deepening/results/L01_同人检出变化全部规格.tsv');t=l[(l.threshold==1)&(l.spec=='all')].groupby('target').first().loc[['Lacticaseibacillus','Lactobacillus']]
ck('纵向信息参与者数',t.people.tolist(),[42,42],'new L01');ck('纵向信息样本数',t.samples.tolist(),[360,362],'new L01')
r1=rd('runs/deepening/results/R01_核心关联全部配套复核.tsv');ck('原分层未通过数',int((~r1.valid).sum()),5,'new R01')
r4=rd('runs/deepening/results/R04_BMI非缺失分层关联.tsv');ck('BMI完整替代子集人数',r4.groupby('spec').n.first().loc[['male_BMI_observed','no_yogurt_BMI_observed']].tolist(),[944,1073],'new R04');ck('BMI完整追加有效模型',int(r4.valid.sum()),12,'new R04')
taxpair=rd('runs/core/表S24g_重新注释15组整体模式比较.tsv');t=taxpair[(taxpair.taxon_A=='Lacticaseibacillus')&(taxpair.taxon_B=='Lactobacillus')].set_index('release').loc[['R07','R06']]
ck('两分类主要配对q',sig(t.q_BH_15),[2.22e-14,8.63e-14],'new C11')
key=rd('runs/core/表S24h_三项突出差异的阈值与分类复核.tsv')
upper('提高阈值及两分类核心差异q上限',key.q_BH_3.max(),8.11e-11,'new C12')
family=json.loads((ROOT/'runs/deepening/results/准备核验.json').read_text(encoding='utf8'))
ck('科汇总菌属数',len(family['family_members']),19,'new family reconstruction')
longindex=rd('data/archive/02_统计输入/冻结统计输入/纵向分析样本_内部索引.tsv')
ck('纵向分析输入样本与人数',[len(longindex),int(longindex.host_subject_id.nunique())],[664,153],'supplied longitudinal selection index; models refit')
cluster=pd.read_csv(ROOT/'runs/sensitivity/within_person/within_person_cluster_results.csv')
t=cluster[(cluster.threshold==1)&(cluster.spec=='all')&(cluster.target=='Lactobacillus')].set_index('background').loc[BG]
ck('纵向聚类OR',rounded(t.OR),[1.632,1.419,1.627],'new N02')
ck('纵向聚类CI低',rounded(t.low),[1.133,1.043,1.137],'new N02')
ck('纵向聚类CI高',rounded(t.high),[2.349,1.930,2.328],'new N02')
ck('纵向聚类q',sig(t.q_BH_6),[.0292,.0535,.0292],'new N02')

coverage=[
 {'section':'1 样本与分析变量','status':'covered','evidence':['preparation_verification.json','source_rebuild_summary.json','diet inventory','Table 1 recalculation here'],'note':'Source tables/BIOM are provided processed AGP inputs; no raw-read processing claim.'},
 {'section':'2 关联模型与多重检验校正','status':'covered','evidence':['core_verification.json: C01-C07','core stacked-model cross-checks','deepening statistical independent checks'],'note':'936 regressions, 15 pairs, 156 omnibus and 2340 comparisons; historical exploratory chronology is documented, not experimentally reverified.'},
 {'section':'3 饮食、联合、检出状态及丰度','status':'covered','evidence':['deepening_validation.json: P/J/S/A01/A02','new PCA scores/loadings','统计独立核验.json'],'note':'Fresh source dietary joins and PCA, full model refits including alternate thresholds and high-read exclusion.'},
 {'section':'4 分类及序列复核','status':'covered_from_supplied_assignments','evidence':['core taxonomy rebuild/validation C08-C15','deepening V01-V03','sensitivity bloom and QC5000 diagnostics'],'note':'VSEARCH/SINTAX annotation itself is not rerun; original R06/R07 classifications are supplied inputs. ASV aggregation and downstream statistical models are regenerated.'},
 {'section':'5 分层、性别、轮次及协变量敏感性','status':'covered','evidence':['deepening R01-R05/H01-H03/M01','sensitivity N01','round and diversity diagnostics'],'note':'Original 5 stratified-model failures and 32-round model failure retained, not converted to successes.'},
 {'section':'6 同人重复采样','status':'covered_from_frozen_longitudinal_index','evidence':['deepening L01','sensitivity N02','within_person diagnostics'],'note':'Refits include 30 original conditional models and 18 clustered uncertainty checks. The 153-person/664-sample longitudinal index is a supplied selection input; no longitudinal cross-target coefficient contrast claimed.'}
]
report={'status':'PASS' if all(c['passed'] for c in checks) else 'FAIL','main_doc_sha256':hashlib.sha256((ROOT/'reference/01_论文正文.docx').read_bytes()).hexdigest(),'supp_doc_sha256':hashlib.sha256((ROOT/'reference/04_补充材料.docx').read_bytes()).hexdigest(),'claim_checks':checks,'table1':{'data_cells_checked':len(table),'all_data_cells_checked':len(table)==14,'footnote_missingness_checked':True,'source':'fresh preparation raw_covariates; all quantiles independently recalculated; percentages denominator 2760'},'supplement_coverage':coverage,'not_claimed':['Fresh raw sequencing processing','Fresh VSEARCH/SINTAX classifications','Independent external cohort validation','Cross-machine execution','Reverification of ethics approvals, bibliography and author identity'],'documents_modified':False}
(ROOT/'reports/manuscript_claims_check.json').write_text(json.dumps(report,ensure_ascii=False,indent=2,default=lambda x:x.item() if isinstance(x,np.generic) else str(x)),encoding='utf8')
print('status',report['status'],'checks',len(checks),'table1 cells',len(table))
for c in checks:
    if not c['passed']:print(json.dumps(c,ensure_ascii=False))
sys.exit(0 if report['status']=='PASS' else 1)
