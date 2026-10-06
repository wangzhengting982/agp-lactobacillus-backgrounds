"""Map all submitted worksheets to archived outputs and producing code.

This checks provenance/cell equality; it does not fit statistical models.
"""
from pathlib import Path
import json, csv, math, hashlib, ast
import openpyxl

ROOT=Path(__file__).resolve().parents[1]
ARC=ROOT/'data/archive'
WB=ROOT/'reference/submitted_results.xlsx'
def rel(p): return str(p.relative_to(ROOT)).replace('\\','/')
def sha(p): return hashlib.sha256(p.read_bytes()).hexdigest()
def eq(a,b):
    if a is None or a=='': return b is None or b=='' or str(b).lower() in ('nan','none','na')
    if isinstance(a,bool): return str(a).lower()==str(b).lower()
    if isinstance(a,(int,float)):
        try: return math.isclose(a,float(b),rel_tol=2e-14,abs_tol=1e-15)
        except (ValueError,TypeError): return False
    return str(a)==str(b)

core='01_核心统计/代码/'
assoc='08_本轮方法复核/provenance/recovered_preparation/new_article/04_复现代码/'
deep='03_饮食与菌群深化/复现代码/'
sex='04_性别与高读数复核/复现代码/'
inputs=['02_统计输入/冻结统计输入/association_covariate_design.tsv','02_统计输入/冻结统计输入/association_target_counts.tsv','02_统计输入/冻结统计输入/association_standardized_clr.tsv']
meta={}
def add(codes,scripts,files=None,notes='',entry_inputs=None):
    for code in codes.split():meta[code]={'producer_scripts':scripts,'archive_output':(files or {}).get(code),'notes':notes,'input_scope':entry_inputs or inputs}
add('C01 C02 C03',[assoc+'association_models/fit_associations.py'],{'C01':'02_统计输入/冻结统计输入/single_genus_all_associations.tsv','C03':'01_核心统计/输入/03_分析与结果/association_models/pc_direct_cross_genus_heterogeneity.json'},'C01 is 936 regressions; C02/C03 use standardized 10-PC scores and joint sandwich inference. Raw preparation and frozen-model-input reproduction are distinct stages.')
add('C04 C07',[core+'01_genus_contrasts.py'],{'C04':'01_核心统计/新增分析/表S22c_全部15组主成分向量比较.tsv','C07':'01_核心统计/新增分析/表S23_整体差异稳健性.tsv'},entry_inputs=inputs+['02_统计输入/冻结统计输入/association_pc_scores.tsv'])
add('C05 C06',[core+'04_complete_background_differences.py'],{'C05':'01_核心统计/新增分析/表S22e_全部156背景属整体差异.tsv','C06':'01_核心统计/新增分析/表S22f_全部2340项背景属系数比较.tsv'})
add('C08 C09 C10',[core+'02_taxonomy_trace.py'],{'C08':'01_核心统计/新增分析/表S24a_重点背景属分类支持与覆盖.tsv','C09':'01_核心统计/新增分析/表S24b_重新注释候选ASV明细.tsv','C10':'01_核心统计/新增分析/表S24c_参考库属标签.tsv'},'Uses archived R06/R07 full-ASV assignments/reference inventories and BIOM counts. Reconstructing these summaries is not fresh SINTAX sequence annotation.',entry_inputs=['01_核心统计/输入/03_分析与结果/taxonomy_validation/','02_统计输入/六份源数据/4168_ASV_table.biom'])
add('C11',[core+'03_taxonomy_PC_contrasts.py'],{'C11':'01_核心统计/新增分析/表S24g_重新注释15组整体模式比较.tsv'},'Original archived implementation uses R06/R07 joint covariance NPZ. The current producer chain refits these arrays before computing contrasts.',entry_inputs=['01_核心统计/输入/03_分析与结果/taxonomy_validation/association_sensitivity/'])
add('C12 C13',[core+'05_key_difference_sensitivity.py'],{'C12':'01_核心统计/新增分析/表S24h_三项突出差异的阈值与分类复核.tsv','C13':'01_核心统计/新增分析/表S24i_三背景两目标属效应量.tsv'},entry_inputs=inputs+['01_核心统计/输入/03_分析与结果/taxonomy_validation/association_sensitivity/'])
add('C14 C15',[assoc+'taxonomy_association_sensitivity.py'],notes='Original electronic output absent as standalone TSV/JSON in archive; numeric targets preserved in the archived core workbook. Archived R06/R07 model matrices permit model refits; full raw-preparation script also expects external annotation completion and raw predictor files.',entry_inputs=['01_核心统计/输入/03_分析与结果/taxonomy_validation/association_sensitivity/'])
for codes,script in [('A01 J01 J02 J03 S01 S02 S03','03_joint_states_abundance.py'),('L01','04_within_person.py'),('R01 R02 R03','06_core_robustness.py'),('R04 R05','08_BMI_strata_resolution.py'),('V01 V02 V03','05_asv.py')]:
    add(codes,[deep+'common.py',deep+script],notes='Original archived implementation and frozen output are retained for provenance. Current migrated producers and execution status are listed separately.')
add('P01 P02',['02_统计输入/饮食预处理/02_diet_core_pretest.py'],{'P01':'02_统计输入/饮食预处理/更广泛饮食调整_全部六配对模型.tsv','P02':'02_统计输入/饮食预处理/更广泛饮食调整_全部属间差异.tsv'},'Refits eight prespecified diet-adjustment specifications; food/nutrient PCA scores must be regenerated for end-to-end reproduction.',entry_inputs=inputs+['02_统计输入/六份源数据/food_exposure_matrix_stool.tsv','02_统计输入/六份源数据/vioscreen_micromacro.tsv'])
add('A02',[sex+'common.py',sex+'08_abundance_influence.py'])
add('H01 H02 H03',[sex+'common.py',sex+'02_sex_contrasts.py'])
add('M01',[sex+'01_intake_audit.py'],{'M01':'04_性别与高读数复核/新增分析结果/M01_额外协变量可用性.json'},'Workbook is a 21-row flattened availability summary of five fields, not an inferential model.',entry_inputs=['02_统计输入/六份源数据/sample_information_10317_matched_4168.tsv','02_统计输入/冻结统计输入/association_covariate_design.tsv'])
add('N01',['08_本轮方法复核/diversity/run_diversity_audit.py','08_本轮方法复核/provenance/02_round_sensitivity.py','08_本轮方法复核/bloom/run_bloom_sensitivity.py','08_本轮方法复核/bloom/run_qc5000_sensitivity.py','07_历史与诊断归档/整理记录/build_supplement_data.py'],notes='Packaging of 24 rows: primary threshold 1; eight fixed scenarios, three backgrounds each. Packaging script does not fit models. Full 32-level round model is invalid and must not be substituted.')
add('N02',['08_本轮方法复核/within_person/run_within_audit.py','07_历史与诊断归档/整理记录/build_supplement_data.py'],{'N02':'08_本轮方法复核/within_person/within_person_cluster_results.csv'},'All 18 participant-clustered conditional-logit models; G/(G-1) correction, t(G-1). Workbook reorders/renames columns and adds explanation fields.')

def read_report(name):
    return json.loads((ROOT/'reports'/name).read_text(encoding='utf-8'))

# Read the actual current validator's output mapping without executing its code.
tree=ast.parse((ROOT/'src/core/06_validate_core.py').read_text(encoding='utf-8'))
core_map=next(ast.literal_eval(node.value) for node in tree.body if isinstance(node,ast.Assign) and any(isinstance(target,ast.Name) and target.id=='MAPPING' for target in node.targets))
core_validation=read_report('core_verification.json')
deep_validation=read_report('deepening_validation.json')
sens_validation=read_report('sensitivity_validation.json')
assert core_validation['status']=='pass' and deep_validation['all_passed'] and sens_validation['all_passed']
current_producers={}
def current(codes,scripts):
    for code in codes.split():current_producers[code]=scripts
current('C01 C05 C06',['src/core/04_complete_background_differences.py'])
current('C02 C03',['src/core/01_genus_contrasts.py','src/core/06_validate_core.py'])
current('C04 C07',['src/core/01_genus_contrasts.py'])
current('C08 C09 C10',['src/core/02_taxonomy_trace.py'])
current('C11',['src/core/02b_refit_taxonomy_pc.py','src/core/03_taxonomy_PC_contrasts.py'])
current('C12 C13',['src/core/05_key_difference_sensitivity.py'])
current('C14',['src/core/02b_refit_taxonomy_pc.py'])
current('C15',['src/core/02b_refit_taxonomy_pc.py','src/core/06_validate_core.py'])
for codes,script in [('A01 J01 J02 J03 S01 S02 S03','03_joint_states_abundance.py'),('L01','04_within_person.py'),('R01 R02 R03','06_core_robustness.py'),('R04 R05','08_BMI_strata_resolution.py'),('V01 V02 V03','05_asv.py'),('A02','10_abundance_influence.py'),('H01 H02 H03','09_sex_contrasts.py')]:
    current(codes,['src/deepening/'+script])
current('P01 P02',['src/deepening/00b_diet_pca.py','scripts/run_deepening.py'])
current('M01',['src/deepening/00c_metadata.py','src/deepening/validate_deepening.py'])
current('N01',['src/sensitivity/diversity/run_diversity_audit.py','src/sensitivity/provenance/02_round_sensitivity.py','src/sensitivity/bloom/run_bloom_sensitivity.py','src/sensitivity/bloom/run_qc5000_sensitivity.py','scripts/validate_sensitivity.py'])
current('N02',['src/sensitivity/within_person/run_within_audit.py','scripts/validate_sensitivity.py'])
assert len(current_producers)==40

def current_result(code,sheet):
    if code in core_map:
        output=ROOT/'runs/core'/core_map[code][0]
        report='reports/core_verification.json'
        checked=next(x for x in core_validation['results'] if x['table']==sheet)
        passed=checked['pass_']
        driver='scripts/run_core.py'
    elif code in ['N01','N02']:
        output=ROOT/'runs/sensitivity'/(code+'_regenerated_submission.tsv')
        report='reports/sensitivity_validation.json'
        checked=next(x for x in sens_validation['checks'] if x['sheet']==sheet)
        passed=checked['passed']
        driver='scripts/run_sensitivity.py; scripts/validate_sensitivity.py'
    else:
        directory='sex_results' if code=='A02' or code.startswith('H') or code=='M01' else 'results'
        output=ROOT/'runs/deepening'/directory/(sheet+'.tsv')
        report='reports/deepening_validation.json'
        checked=next(x for x in deep_validation['checks'] if x['table']==sheet)
        passed=checked['passed']
        driver='scripts/run_deepening.py'
    assert output.is_file(),output
    assert all((ROOT/p).is_file() for p in current_producers[code]),(code,current_producers[code])
    assert passed,(code,checked)
    return {'current_producer_script':current_producers[code],
            'current_pipeline_entry':driver,
            'current_output':rel(output),
            'current_output_sha256':sha(output),
            'current_verification_report':report,
            'current_status':'PASS',
            'current_max_absolute_error':checked.get('max_absolute_error'),
            'current_scope':'Recomputed locally and compared with the submitted worksheet. Supplied processed source tables and taxonomy labels are inputs; FASTQ processing and SINTAX classification were not rerun.'}

work=openpyxl.load_workbook(WB,data_only=True)
rows=[]
for item in list(work.worksheets[0].values)[1:]:
    code,title,sheet,section,n,oldwb,oldsheet=item
    m=meta[code]
    if not m['archive_output'] and code[0] in ('A','J','L','R','S','V','H'):
        hits=list(ARC.glob('03_饮食与菌群深化/完整分析结果/'+code+'_*.tsv'))+list(ARC.glob('04_性别与高读数复核/新增分析结果/'+code+'_*.tsv'))
        assert len(hits)==1,(code,hits)
        m['archive_output']=str(hits[0].relative_to(ARC)).replace('\\','/')
    prior=ARC/'10_本轮精简与历史材料/原补充数据'/oldwb
    old=openpyxl.load_workbook(prior,data_only=True)
    vals=list(work[sheet].values);previous=list(old[oldsheet].values)
    old_equal=len(vals)==len(previous) and all(len(a)==len(b) and all(eq(x,y) for x,y in zip(a,b)) for a,b in zip(vals,previous))
    out=ARC/m['archive_output'] if m['archive_output'] else None
    direct_check=None
    if out and out.suffix in ('.tsv','.csv'):
        with out.open(encoding='utf-8-sig',newline='') as f: data=list(csv.reader(f,delimiter='\t' if out.suffix=='.tsv' else ','))
        direct_check={'equal_rows':len(vals)==len(data),'equal_cells':len(vals)==len(data) and all(len(a)==len(b) and all(eq(x,y) for x,y in zip(a,b)) for a,b in zip(vals,data))}
        if code=='N02':direct_check['reason']='Workbook uses translated headers/reordered columns/added explanatory columns; source imported by packaging script.'
    sources=m['producer_scripts']
    assert all((ARC/p).is_file() for p in sources),(code,sources)
    if out: assert out.exists(),out
    row={'code':code,'worksheet':sheet,'rows':n,'columns':work[sheet].max_column,'section':section,'archived_workbook':rel(prior),'archived_worksheet':oldsheet,'archived_workbook_cells_match':old_equal,'archived_producer_scripts':[rel(ARC/p) for p in sources],'archive_output':rel(out) if out else None,'archive_output_sha256':sha(out) if out else None,'direct_output_check':direct_check,'archived_input_scope':[rel(ARC/p) for p in m['input_scope']],'archive_provenance_notes':m['notes']}
    row.update(current_result(code,sheet))
    assert old_equal,(code,'archived original mismatch')
    rows.append(row)
assert len(rows)==40 and len(meta)==40
report={'submitted_workbook_sha256':sha(WB),'number_of_result_sheets':40,'numeric_comparison_tolerance':{'relative':2e-14,'absolute':1e-15,'reason':'This tolerance is for archive-to-workbook provenance comparison. Current model-run tolerances are defined separately in each current_verification_report.'},'all_original_workbook_cells_match':all(r['archived_workbook_cells_match'] for r in rows),'all_producer_scripts_found':True,'current_validated_sheets':sum(r['current_status']=='PASS' for r in rows),'current_all_passed':all(r['current_status']=='PASS' for r in rows),'tables':rows,'excluded_historical_scope':'The 40 retained submitted sheets are mapped and their current validation reports pass. Unsubmitted historical prediction, daily-cohort and external-availability branches are outside this current 40-sheet claim.','validation_scope':'All 40 submitted result worksheets were regenerated and validated locally from supplied processed source tables and annotation inputs. No FASTQ processing or new SINTAX classification is claimed.'}
(ROOT/'reports/table_map.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
with (ROOT/'reports/table_map.csv').open('w',encoding='utf-8-sig',newline='') as f:
    wr=csv.DictWriter(f,fieldnames=['code','worksheet','rows','columns','current_producer_script','current_pipeline_entry','current_output','current_verification_report','current_status','current_max_absolute_error','archive_output','archived_producer_scripts','archived_workbook_cells_match','archive_provenance_notes'])
    wr.writeheader()
    for r in rows:wr.writerow({k:'; '.join(r[k]) if isinstance(r[k],list) else r[k] for k in wr.fieldnames})
print(json.dumps({'sheets':len(rows),'old_workbook_matches':sum(r['archived_workbook_cells_match'] for r in rows),'direct_tabular_matches':sum(bool(r['direct_output_check'] and r['direct_output_check']['equal_cells']) for r in rows),'transformed_or_json_sources':[r['code'] for r in rows if not r['direct_output_check'] or not r['direct_output_check']['equal_cells']]},indent=2))
